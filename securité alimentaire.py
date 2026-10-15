from datetime import datetime
from io import BytesIO
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Audit GFS - Sécurité Alimentaire",
    page_icon="🍽️",
    layout="wide",
)

st.title("🍽️ Audit Global Food Safety (GFS) - Marriott")
st.markdown(
    "**Plateforme de contrôle d'hygiène et de sécurité des aliments**"
)

# Initialisation des réponses dans la session
if "answers" not in st.session_state:
    st.session_state.answers = {}

# ---------------------------------------------------------
# DÉFINITION COMPLÈTE DES 43 POINTS DE CONTRÔLE DU PDF GFS
# ---------------------------------------------------------
GFS_CHECKLIST = {
    "FS-Documentation": [
        (
            "FSDOC001",
            "1. Les procédures de notification des intoxications alimentaires"
            " de l'année en cours sont documentées et accessibles.",
            2,
        ),
        (
            "FSDOC002",
            "2. Le document attestant que les employés qui manipulent des"
            " aliments ont été formés à la sécurité sanitaire des aliments est"
            " en cours de validité.",
            2,
        ),
        (
            "FSDOC003",
            "3. Les responsables ont reçu une certification après avoir suivi un"
            " programme approuvé.",
            2,
        ),
        (
            "FSDOC004",
            "4. Un programme de lutte contre les nuisibles a été adopté et mis"
            " en place (pas de présence de nuisibles).",
            2,
        ),
        (
            "FSDOC005",
            "5. Les formulaires d'enregistrement des températures du"
            " réfrigérateur/congélateur (HACCP A-3) sont renseignés et"
            " conservés pendant 90 jours.",
            2,
        ),
        (
            "FSDOC006",
            "6. Les formulaires d'enregistrement des températures (HACCP A-1)"
            " pour la cuisson, la conservation et le réchauffage sont"
            " conservés pendant 90 jours.",
            16,
        ),
        (
            "FSDOC007",
            "7. Les formulaires d'enregistrement du refroidissement des"
            " aliments (HACCP A-2) sont renseignés et conservés pendant 90"
            " jours.",
            16,
        ),
        (
            "FSDOC009",
            "8. Les fiches de suivi du nettoyage des postes de travail sont"
            " renseignées et conservées pendant 90 jours.",
            16,
        ),
    ],
    "FS-Facilities (Locaux)": [
        (
            "FSFAC001",
            "9. Les systèmes d'aération et hottes de ventilation de la cuisine"
            " sont nettoyés et entretenus par des professionnels (tous les 6"
            " mois).",
            2,
        ),
        (
            "CABFS102",
            "10. Les machines à glaçons sont propres et en bon état (sans"
            " moisissure, rouille ou saleté).",
            2,
        ),
        (
            "FSFAC003",
            "11. Les sols, les murs et les plafonds sont propres et ne"
            " présentent pas de traces de poussière, de débris ou d'eau"
            " stagnante.",
            2,
        ),
        (
            "FSFAC004",
            "12. Les sols, les murs et les plafonds sont facilement"
            " nettoyables et en bon état.",
            2,
        ),
        (
            "FSFAC005",
            "13. Le système de ventilation est adapté (bouches d'aération,"
            " grilles et filtres propres).",
            2,
        ),
        (
            "FSFAC006",
            "14. Les espaces de livraison et d'entrée sont protégés contre les"
            " insectes et les rongeurs.",
            2,
        ),
        (
            "FSFAC007",
            "15. Les espaces réservés aux poubelles sont toujours propres et"
            " bien organisés.",
            2,
        ),
        (
            "FSFAC008",
            "16. Les espaces réservés aux employés (cafétéria/cantine/salle de"
            " pause) sont propres et bien entretenus.",
            4,
        ),
        (
            "FSFAC009",
            "17. Un espace est dédié au lavage des mains dans les aires de"
            " manipulation des aliments (eau chaude à 38°C, savon, essuie-mains"
            " jetable).",
            2,
        ),
        (
            "CABFS105",
            "18. Les lave-vaisselle sont correctement entretenus et utilisés"
            " (désinfection à 71°C ou dosage chimique correct).",
            4,
        ),
        (
            "FSFAC011",
            "19. Les produits chimiques sont utilisés et stockés selon les"
            " normes (étiquetage, séparation des aliments).",
            2,
        ),
        (
            "FSFAC012",
            "20. Les bacs de plonge et de désinfection de l'équipement et des"
            " ustensiles sont correctement entretenus.",
            2,
        ),
        (
            "FSFAC013",
            "21. Les aliments et emballages alimentaires sont correctement"
            " stockés à au moins 15 cm du sol.",
            2,
        ),
        (
            "FSFAC014",
            "22. La méthode FIFO (Premier Entré - Premier Sorti) est appliquée"
            " dans les aires de production et de stockage.",
            2,
        ),
        (
            "FSFAC015",
            "23. La température des réfrigérateurs (<= 5°C) et congélateurs"
            " (<= -15°C) respecte la norme.",
            16,
        ),
        (
            "CABFS103",
            "24. Les aliments froids à risque sont maintenus à une température"
            " <= 5°C dans tous les systèmes de conservation.",
            16,
        ),
        (
            "FSFAC017",
            "25. Les aliments chauds à risque sont maintenus à une température"
            " >= 57°C dans tous les systèmes de conservation.",
            16,
        ),
    ],
    "FS-Food Handling (Manipulation des aliments)": [
        (
            "FSFH001",
            "26. Les employés qui manipulent les aliments connaissent les"
            " températures de cuisson standard (ex: volaille 74°C, viande hachée"
            " 68°C).",
            2,
        ),
        (
            "FSFH015",
            "27. Les steaks hachés frais ou surgelés respectent les normes"
            " microbiologiques et de cuisson à cœur.",
            4,
        ),
        (
            "FSFH003",
            "28. Les aliments proviennent de fournisseurs commerciaux agréés"
            " (étiquettes de fruits de mer conservées 90 jours).",
            2,
        ),
        (
            "FSFH004",
            "29. La date limite de conservation (max 7 jours) est indiquée au"
            " moment de la préparation sur les aliments à risque prêts à"
            " consommer.",
            2,
        ),
        (
            "FSFH005",
            "30. Les produits alimentaires ne sont pas conservés ou servis"
            " au-delà de leur date limite de consommation.",
            2,
        ),
        (
            "FSFH006",
            "31. Les aliments et surfaces sont protégés des dangers"
            " microbiologiques, physiques et chimiques (pas de viandes crues"
            " au-dessus des cuits).",
            4,
        ),
        (
            "FSFH007",
            "32. Les aliments à risque sont décongelés selon les règles"
            " (au froid <= 5°C ou sous eau courante <= 21°C).",
            2,
        ),
        (
            "FSFH008",
            "33. Les ustensiles en cours d'utilisation (y compris pelles à"
            " glace) sont manipulés et stockés de manière hygiénique.",
            2,
        ),
        (
            "CABFS106",
            "34. Les directives relatives aux planches à découper respectent"
            " le code couleur et les normes.",
            4,
        ),
        (
            "FSFH010",
            "35. Les surfaces en contact avec les aliments (équipement/ustensiles)"
            " sont durables, non toxiques, nettoyables et en bon état.",
            2,
        ),
        (
            "FSFH011",
            "36. L'équipement et les ustensiles sont correctement nettoyés"
            " (joints, poignées, étagères, rails).",
            2,
        ),
        (
            "FSFH012",
            "37. Les surfaces qui ne sont pas en contact direct avec les"
            " aliments sont propres et en bon état.",
            2,
        ),
        (
            "FSFH013",
            "38. Les solutions de désinfection chimique sont utilisées à la"
            " bonne concentration et température.",
            4,
        ),
        (
            "FSFH014",
            "39. Des kits de test (bandelettes) pour désinfectant sont"
            " facilement accessibles.",
            4,
        ),
    ],
    "FS-Personal Hygiene (Hygiène personnelle)": [
        (
            "FSPH001",
            "40. Les employés se lavent fréquemment et correctement les"
            " mains aux moments clés.",
            16,
        ),
        (
            "FSPH002",
            "41. Aucun contact à mains nues n'a été constaté avec les aliments"
            " prêts à consommer.",
            4,
        ),
        (
            "FSPH003",
            "42. Les employés présentant des symptômes contagieux sont exclus/restreints"
            " et les consignes sont affichées.",
            4,
        ),
        (
            "FSPH004",
            "43. Il est interdit de boire, manger ou fumer dans les aires où"
            " sont manipulés les aliments.",
            2,
        ),
    ],
}

# Barre latérale : Navigation simple
st.sidebar.header("📋 Navigation Audit")
selected_category = st.sidebar.radio(
    "Choisir la catégorie :", list(GFS_CHECKLIST.keys())
)

# Formulaire principal
st.header(f"Section : {selected_category}")
st.write("Évaluez chaque point de contrôle ci-dessous :")

for code, consigne, pts in GFS_CHECKLIST[selected_category]:
    with st.container():
        col1, col2, col3 = st.columns([4, 2, 3])
        with col1:
            st.markdown(f"**[{code}]** {consigne}")
            st.caption(f"Pondération : **{pts} Pts**")
        with col2:
            key_status = f"status_{code}"
            current_status = st.session_state.answers.get(
                code, {}
            ).get("statut", "Oui")
            status = st.radio(
                "Statut",
                ["Oui", "Non", "N/A"],
                index=["Oui", "Non", "N/A"].index(current_status),
                key=key_status,
                horizontal=True,
            )
        with col3:
            key_comm = f"comm_{code}"
            current_comm = st.session_state.answers.get(
                code, {}
            ).get("commentaire", "")
            commentaire = st.text_input(
                "Observations / Action corrective",
                value=current_comm,
                key=key_comm,
                placeholder="Ex: Température relevée 7.5°C...",
            )

        # Mettre à jour l'état du contrôle dans session_state
        st.session_state.answers[code] = {
            "Catégorie": selected_category,
            "Code": code,
            "Consigne": consigne,
            "Points": pts,
            "Statut": status,
            "Commentaire": commentaire,
        }
        st.markdown("---")

# ---------------------------------------------------------
# FONCTION DE GÉNÉRATION EXCEL COLORÉ ET DESIGN
# ---------------------------------------------------------
def generate_styled_excel(data_dict):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rapport Audit GFS"
    ws.views.sheetView[0].showGridLines = True

    # Couleurs et Styles
    HEADER_BG = "1F4E78"  # Bleu Foncé Elegant
    HEADER_FG = "FFFFFF"  # Blanc
    ZEBRA_BG = "F2F5F9"  # Gris-Bleu très clair

    font_title = Font(name="Segoe UI", size=16, bold=True, color="1F4E78")
    font_header = Font(name="Segoe UI", size=11, bold=True, color=HEADER_FG)
    font_data = Font(name="Segoe UI", size=10)

    fill_header = PatternFill(
        start_color=HEADER_BG, end_color=HEADER_BG, fill_type="solid"
    )
    fill_zebra = PatternFill(
        start_color=ZEBRA_BG, end_color=ZEBRA_BG, fill_type="solid"
    )

    # Couleurs de Statut
    fill_oui = PatternFill(
        start_color="D9EAD3", end_color="D9EAD3", fill_type="solid"
    )  # Vert clair
    fill_non = PatternFill(
        start_color="FCE5CD", end_color="FCE5CD", fill_type="solid"
    )  # Orange/Rouge clair
    fill_na = PatternFill(
        start_color="EFEFEF", end_color="EFEFEF", fill_type="solid"
    )  # Gris clair

    font_oui = Font(name="Segoe UI", size=10, bold=True, color="274E13")
    font_non = Font(name="Segoe UI", size=10, bold=True, color="783F04")
    font_na = Font(name="Segoe UI", size=10, bold=True, color="666666")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Titre du Rapport
    ws.merge_cells("A1:F1")
    ws["A1"] = "RAPPORT D'AUDIT GLOBAL FOOD SAFETY (GFS)"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center")

    ws["A2"] = f"Date de l'audit : {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="595959")

    # En-têtes du Tableau
    headers = [
        "Catégorie",
        "Code",
        "Consigne de Contrôle",
        "Pts",
        "Statut",
        "Observations / Actions Correctives",
    ]
    ws.append([])  # Ligne 3 vide

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    # Remplissage des données
    row_idx = 5
    for cat, items in GFS_CHECKLIST.items():
        for code, consigne, pts in items:
            ans = data_dict.get(
                code,
                {
                    "Statut": "Non évalué",
                    "Commentaire": "",
                },
            )
            statut = ans.get("Statut", "Oui")
            commentaire = ans.get("Commentaire", "")

            c_cat = ws.cell(row=row_idx, column=1, value=cat)
            c_code = ws.cell(row=row_idx, column=2, value=code)
            c_cons = ws.cell(row=row_idx, column=3, value=consigne)
            c_pts = ws.cell(row=row_idx, column=4, value=pts)
            c_stat = ws.cell(row=row_idx, column=5, value=statut)
            c_comm = ws.cell(row=row_idx, column=6, value=commentaire)

            for cell in [c_cat, c_code, c_cons, c_pts, c_stat, c_comm]:
                cell.font = font_data
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
                if row_idx % 2 == 0:
                    cell.fill = fill_zebra

            c_code.alignment = Alignment(horizontal="center", vertical="center")
            c_pts.alignment = Alignment(horizontal="center", vertical="center")
            c_stat.alignment = Alignment(horizontal="center", vertical="center")

            if statut == "Oui":
                c_stat.fill = fill_oui
                c_stat.font = font_oui
            elif statut == "Non":
                c_stat.fill = fill_non
                c_stat.font = font_non
            elif statut == "N/A":
                c_stat.fill = fill_na
                c_stat.font = font_na

            row_idx += 1

    column_widths = {
        "A": 25,
        "B": 15,
        "C": 55,
        "D": 8,
        "E": 12,
        "F": 45,
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


# ---------------------------------------------------------
# BOUTON DE TÉLÉCHARGEMENT DIRECT SUR LA PAGE PRINCIPALE
# ---------------------------------------------------------
st.markdown("---")
st.subheader("📥 Exportation du Rapport d'Audit")
st.write(
    "Cliquez sur le bouton ci-dessous pour télécharger votre rapport Excel"
    " mis en forme et coloré :"
)

excel_data = generate_styled_excel(st.session_state.answers)

st.download_button(
    label="📊 Télécharger le Rapport Excel Coloré",
    data=excel_data,
    file_name=f"Rapport_Audit_GFS_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
